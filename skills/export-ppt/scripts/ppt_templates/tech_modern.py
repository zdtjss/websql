"""
科技现代模板
蓝色系渐变 + 青色强调 + 橙色点缀，适合互联网/科技企业场景。
"""

from pptx.util import Pt
from pptx.enum.text import PP_ALIGN

from .base import PPTTemplate


class TechModernTemplate(PPTTemplate):
    name = "tech_modern"
    label = "科技现代"

    def build_cover(self, slide, title, subtitle, presenter=""):
        from shared.utils import clean_surrogates, format_date_cn
        self._gradient_bg(slide,
                           self.scheme["primary_dark"],
                           self.scheme["primary"],
                           self.scheme["primary_light"])
        self._accent_bar(slide, 2.6, 0.05, self.scheme["accent"])

        self._textbox(slide, 1.8, 1.6, 9.5, 2.4, clean_surrogates(title),
                      size=46, bold=True, color=self.scheme["bg_white"])
        self._textbox(slide, 1.8, 3.6, 9.5, 1.2, clean_surrogates(subtitle),
                      size=20, color=self.scheme["accent_light"])

        now = format_date_cn()
        info = f"▸  {now}"
        if presenter:
            info += f"  |  汇报人：{clean_surrogates(presenter)}"
        info += f"  |  {self.config.brand_name}"
        self._textbox(slide, 1.8, 4.8, 9.5, 0.5, info, size=12, color=self.scheme["text_muted"])
        self._textbox(slide, 1.8, 6.3, 9.5, 0.4, "Powered by WebSQL AI",
                      size=10, color=self.scheme["text_muted"])

    def build_toc(self, slide, sections, slide_num, total):
        from shared.utils import clean_surrogates
        self._solid_bg(slide, self.scheme["bg_slide"])
        self._accent_bar(slide, 0, 0.05, self.scheme["accent"])
        self._textbox(slide, 1.2, 0.3, 10.0, 0.9, "目  录",
                      size=32, bold=True, color=self.scheme["primary"])

        def draw_item(left, top, num, title_text, desc_text=""):
            circle = slide.shapes.add_shape(9, Inches(left), Inches(top),
                                           Inches(0.5), Inches(0.5))
            circle.fill.solid()
            circle.fill.fore_color.rgb = self._rgb(self.scheme["accent"])
            circle.line.fill.background()
            tf = circle.text_frame
            tf.word_wrap = False
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            r = p.add_run()
            r.text = str(num)
            self._font(r, size=16, bold=True, color=self.scheme["bg_white"], font_cn="Calibri")

            self._textbox(slide, left + 0.8, top - 0.02, 8.0, 0.4, clean_surrogates(title_text),
                          size=18, bold=True, color=self.scheme["text_primary"])
            if desc_text:
                self._textbox(slide, left + 0.8, top + 0.32, 8.0, 0.3, clean_surrogates(desc_text),
                              size=10, color=self.scheme["text_secondary"])

        y = 1.5
        for i, sec in enumerate(sections[:10]):
            draw_item(2.0, y, i + 1, sec["title"], sec.get("desc", ""))
            y += 0.85

        self._page_num(slide, slide_num, total)

    def build_section(self, slide, num, title, slide_num, total):
        from shared.utils import num_to_cn, clean_surrogates
        self._gradient_bg(slide,
                           self.scheme["primary_dark"],
                           self.scheme["primary"],
                           self.scheme["primary_light"])
        self._textbox(slide, 1.5, 1.8, 3.0, 1.2, f"PART {num:02d}",
                      size=44, bold=True, color=self.scheme["accent"], font_cn="Calibri")
        self._textbox(slide, 1.5, 3.4, 10.0, 0.8, f"第{num_to_cn(num)}部分",
                      size=18, color=self.scheme["text_muted"])
        self._textbox(slide, 1.5, 3.9, 10.0, 1.0, clean_surrogates(title),
                      size=34, bold=True, color=self.scheme["bg_white"])
        self._page_num(slide, slide_num, total)

    def build_content(self, slide, title, content_text, slide_num, total):
        from shared.utils import clean_surrogates
        self._solid_bg(slide, self.scheme["bg_slide"])
        self._accent_bar(slide, 0, 0.05, self.scheme["accent"])
        self._textbox(slide, 1.2, 0.3, 10.0, 0.8, clean_surrogates(title),
                      size=28, bold=True, color=self.scheme["primary"])
        self._textbox(slide, 1.2, 1.4, 10.5, 5.3, clean_surrogates(content_text),
                      size=15, color=self.scheme["text_primary"])
        self._page_num(slide, slide_num, total)

    def build_ending(self, slide, title, slide_num, total):
        from shared.utils import clean_surrogates, format_date_cn
        self._gradient_bg(slide,
                           self.scheme["primary_dark"],
                           self.scheme["primary"],
                           self.scheme["primary_light"])
        self._textbox(slide, 2.0, 2.4, 9.5, 1.5, "Thank You",
                      size=52, bold=True, color=self.scheme["bg_white"], align=PP_ALIGN.CENTER)
        self._textbox(slide, 2.0, 3.6, 9.5, 0.5, "谢谢观看",
                      size=28, bold=True, color=self.scheme["accent_light"], align=PP_ALIGN.CENTER)
        self._accent_bar(slide, 4.3, 0.03, self.scheme["accent"])
        self._textbox(slide, 2.0, 4.8, 9.5, 0.6, clean_surrogates(title),
                      size=18, color=self.scheme["text_muted"], align=PP_ALIGN.CENTER)
        self._textbox(slide, 2.0, 5.5, 9.5, 0.5,
                      f"{self.config.brand_name} · {format_date_cn()}",
                      size=13, color=self.scheme["text_muted"], align=PP_ALIGN.CENTER)
        self._page_num(slide, slide_num, total)
