#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""WebSQL 上传文件内容解析脚本（第一方受信代码，仅读不改）。

契约（与 internal/ai/agent/fileparse/python.go 严格对齐）：
  stdin : 单行 JSON  {"path": str, "kind": "docx"|"pdf", "maxChars": int}
  stdout: 单行 JSON  {"ok": true,  "text": str, "totalChars": int, "pages": int,
                      "meta": {...}, "notice": str}
                    或 {"ok": false, "error": str}
  stderr: 仅日志（stdout 必须只有一行 JSON，否则 Go 侧无法解析）

安全约定：
  - 只读取 stdin 指定的单个文件，不创建/修改任何文件，不访问网络；
  - 不执行文档中的任何指令（宏、外链、嵌入对象一律忽略）。
"""

import json
import os
import sys

MAX_ERROR_LEN = 500


def emit(payload):
    """输出结果：stdout 只允许一行 JSON。"""
    sys.stdout.write(json.dumps(payload, ensure_ascii=False))
    sys.stdout.flush()


def fail(message):
    emit({"ok": False, "error": message})


def friendlify(exc):
    """把解析库异常转为可展示的中文提示（不暴露服务器路径与堆栈）。"""
    name = type(exc).__name__
    text = str(exc)[:MAX_ERROR_LEN].replace("\n", " ")
    lowered = (name + " " + text).lower()
    if "password" in lowered or "encrypted" in lowered:
        return "PDF 已加密，请解密后重新上传"
    if ("packagenotfound" in lowered or "zipfile" in lowered or "badzip" in lowered
            or "keyerror" in lowered or "no item named" in lowered):
        return "文件损坏或格式不正确，请确认后重新上传"
    if "modulenotfound" in lowered or "importerror" in lowered:
        return "解析组件不可用：" + text
    return "解析失败（%s）：%s" % (name, text)


def parse_docx(path):
    """python-docx 提取正文：段落 + 表格（表格每行以制表符分隔）。

    说明：只覆盖文档主体（document.paragraphs / document.tables），
    页眉页脚、文本框、批注不在提取范围内。
    """
    import docx  # python-docx

    document = docx.Document(path)
    blocks = []
    for para in document.paragraphs:
        text = (para.text or "").strip()
        if text:
            blocks.append(text)
    for table in document.tables:
        for row in table.rows:
            cells = [(cell.text or "").strip().replace("\n", " ") for cell in row.cells]
            if any(cells):
                blocks.append("\t".join(cells))
    meta = {"paragraphs": len(document.paragraphs), "tables": len(document.tables)}
    return "\n".join(blocks), 0, meta, ""


def parse_pdf(path):
    """pdfplumber 逐页提取文本（带页码标记，便于模型引用页码）。"""
    import pdfplumber

    blocks = []
    empty_pages = 0
    pages = 0
    with pdfplumber.open(path) as pdf:
        pages = len(pdf.pages)
        for index, page in enumerate(pdf.pages, start=1):
            try:
                text = (page.extract_text() or "").strip()
            except Exception:
                # 单页抽取失败不应中断整份文档
                text = ""
            if not text:
                empty_pages += 1
                continue
            blocks.append("----- 第 %d 页 -----\n%s" % (index, text))

    meta = {"emptyPages": empty_pages}
    if not blocks:
        notice = "未能从 PDF 中提取到文本，可能是扫描件或纯图片型 PDF（可改为上传页面截图，由视觉模型识别）"
    elif empty_pages:
        notice = "有 %d 页未提取到文本（可能为图片页）" % empty_pages
    else:
        notice = ""
    return "\n\n".join(blocks), pages, meta, notice


def main():
    try:
        raw = sys.stdin.read()
        request = json.loads(raw) if raw.strip() else {}
    except Exception:
        return fail("解析请求格式错误")

    path = str(request.get("path") or "")
    kind = str(request.get("kind") or "").lower()
    try:
        max_chars = int(request.get("maxChars") or 0)
    except Exception:
        max_chars = 0

    if not path or not os.path.isfile(path):
        return fail("文件不存在或不可读（可能已过期，请重新上传）")

    try:
        if kind == "docx":
            text, pages, meta, notice = parse_docx(path)
        elif kind == "pdf":
            text, pages, meta, notice = parse_pdf(path)
        else:
            return fail("不支持的文件类型：%s" % (kind or "未知"))
    except Exception as exc:  # noqa: BLE001 —— 统一转为结构化错误返回
        return fail(friendlify(exc))

    total = len(text)
    if max_chars > 0 and total > max_chars:
        text = text[:max_chars]
        notice = (notice + "；" if notice else "") + ("文档过长，仅保留前 %d 字符" % max_chars)

    emit({
        "ok": True,
        "text": text,
        "totalChars": total,
        "pages": pages,
        "meta": meta,
        "notice": notice,
    })


if __name__ == "__main__":
    main()
